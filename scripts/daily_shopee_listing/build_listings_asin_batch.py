"""ASIN直接投入モード: 18ASIN(2件除外)のShopee出品ファイル生成スクリプト"""
import sys
import os

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '.claude', 'skills', 'shopee-listing-upload'))

from xlsx_fix import load_fixed_copy, convert_to_shared_strings
from descriptions import DESCRIPTIONS
import openpyxl
from pricing_calc import calc_local_price

DATE_STR = "2026-09-18"
FRAME_BASE_URL = "https://raw.githubusercontent.com/jyankel0726switch-code/ai-management-shopee/claude/admiring-bardeen-itpmle/framed-images"

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), '..', '..', '.claude', 'skills', 'shopee-listing-upload', 'template')
OUT_DIR = os.path.dirname(__file__)

DELIVERY_COLS = {"SG": [34, 35, 36], "MY": [34, 35], "TH": [34], "PH": [34]}
LAST_ITEM_IMAGE_COL = 29  # Item Image 8

# 各行: category, title(先頭行のみ), description_key(先頭行のみ), parent_sku, sku,
# variation_integration_no, variation_name1, option1, cost_jpy, stock, jan,
# weight_kg, length, width, height, images(親情報流用時は共通リストを渡す)
ROWS = [
    # P1: S.H.Figuarts Tanjiro (variation: Style x3)
    dict(category="101385", title="P1", desc="P1", parent_sku="B0FFS31TYR", sku="B0FFS31TYR-01",
         vint="VG-B0FFS31TYR", vname1="Style", opt1="Infinity Castle Ver.",
         cost_jpy=5800, stock=10, jan="0000000000000",
         weight_kg=0.2, length=15, width=5, height=14,
         images=["https://m.media-amazon.com/images/I/51iT7MjwyHL.jpg",
                 "https://m.media-amazon.com/images/I/31n6i6qs0XL.jpg",
                 "https://m.media-amazon.com/images/I/513ACDOmZyL.jpg",
                 "https://m.media-amazon.com/images/I/512WkOtw4kL.jpg",
                 "https://m.media-amazon.com/images/I/4188R8xEKKL.jpg",
                 "https://m.media-amazon.com/images/I/51+lk2moUFL.jpg",
                 "https://m.media-amazon.com/images/I/514hV5fk46L.jpg"]),
    dict(category="101385", title=None, desc=None, parent_sku="B0FFS31TYR", sku="B0FFS31TYR-02",
         vint="VG-B0FFS31TYR", vname1="Style", opt1="Standard Ver.",
         cost_jpy=5800, stock=10, jan="0000000000000",
         weight_kg=0.2, length=15, width=5, height=14,
         images=["https://m.media-amazon.com/images/I/51iT7MjwyHL.jpg",
                 "https://m.media-amazon.com/images/I/31n6i6qs0XL.jpg",
                 "https://m.media-amazon.com/images/I/513ACDOmZyL.jpg",
                 "https://m.media-amazon.com/images/I/512WkOtw4kL.jpg",
                 "https://m.media-amazon.com/images/I/4188R8xEKKL.jpg",
                 "https://m.media-amazon.com/images/I/51+lk2moUFL.jpg",
                 "https://m.media-amazon.com/images/I/514hV5fk46L.jpg"]),
    dict(category="101385", title=None, desc=None, parent_sku="B0FFS31TYR", sku="B0FFS31TYR-03",
         vint="VG-B0FFS31TYR", vname1="Style", opt1="Soul EFFECT Water Blue Ver.",
         cost_jpy=5800, stock=10, jan="0000000000000",
         weight_kg=0.2, length=15, width=5, height=14,
         images=["https://m.media-amazon.com/images/I/51iT7MjwyHL.jpg",
                 "https://m.media-amazon.com/images/I/31n6i6qs0XL.jpg",
                 "https://m.media-amazon.com/images/I/513ACDOmZyL.jpg",
                 "https://m.media-amazon.com/images/I/512WkOtw4kL.jpg",
                 "https://m.media-amazon.com/images/I/4188R8xEKKL.jpg",
                 "https://m.media-amazon.com/images/I/51+lk2moUFL.jpg",
                 "https://m.media-amazon.com/images/I/514hV5fk46L.jpg"]),

    # P2: Demon Slayer Vol.1 Blu-ray (variation: Format x2, individual prices known)
    dict(category="100740", title="P2", desc="P2", parent_sku="B07QS7YTWS", sku="B07QS7YTWS",
         vint="VG-B07QS7YTWS", vname1="Format", opt1="Blu-ray",
         cost_jpy=4880, stock=1, jan="4534530117564",
         weight_kg=0.19, length=30, width=10, height=20,
         images=["https://m.media-amazon.com/images/I/81KRtGoGc-L.jpg",
                 "https://m.media-amazon.com/images/I/613lHm0DTBL.jpg"]),
    dict(category="100740", title=None, desc=None, parent_sku="B07QS7YTWS", sku="B07QS7YTWS-DVD",
         vint="VG-B07QS7YTWS", vname1="Format", opt1="DVD",
         cost_jpy=4438, stock=10, jan="0000000000000",
         weight_kg=0.19, length=30, width=10, height=20,
         images=["https://m.media-amazon.com/images/I/81KRtGoGc-L.jpg",
                 "https://m.media-amazon.com/images/I/613lHm0DTBL.jpg"]),

    # P3-P16: 単品(バリエーションなし)
    dict(category="101392", title="P3", desc="P3", parent_sku="B0H5K9HV63", sku="B0H5K9HV63",
         cost_jpy=1764, stock=5, jan="0000000000000",
         weight_kg=0.15, length=10, width=10, height=11,
         images=["https://m.media-amazon.com/images/I/61wmYy4v1CL.jpg",
                 "https://m.media-amazon.com/images/I/31Xs1eGeyuL.jpg"]),
    dict(category="101385", title="P4", desc="P4", parent_sku="B0FG794W4J", sku="B0FG794W4J",
         cost_jpy=5030, stock=5, jan="0000000000000",
         weight_kg=0.3, length=8, width=8, height=11,
         images=["https://m.media-amazon.com/images/I/41FNBaRUTVL.jpg",
                 "https://m.media-amazon.com/images/I/41HvH8e+B0L.jpg",
                 "https://m.media-amazon.com/images/I/411Dc8gbniL.jpg",
                 "https://m.media-amazon.com/images/I/41RMrh9KeIL.jpg",
                 "https://m.media-amazon.com/images/I/41QXNJE00YL.jpg",
                 "https://m.media-amazon.com/images/I/41n+bizrKDL.jpg",
                 "https://m.media-amazon.com/images/I/51TcKB6euFL.jpg"]),
    dict(category="101392", title="P5", desc="P5", parent_sku="B0H1YN5J4V", sku="B0H1YN5J4V",
         cost_jpy=3500, stock=2, jan="0000000000000",
         weight_kg=0.299, length=18.7, width=15.2, height=11.4,
         images=["https://m.media-amazon.com/images/I/41L7q8u6DhL.jpg",
                 "https://m.media-amazon.com/images/I/41aAgjZ6-CL.jpg",
                 "https://m.media-amazon.com/images/I/411+EhzyxML.jpg",
                 "https://m.media-amazon.com/images/I/41-9oogmVGL.jpg",
                 "https://m.media-amazon.com/images/I/715XYiPIY1L.jpg"]),
    dict(category="101399", title="P6", desc="P6", parent_sku="B0H2KXG1RW", sku="B0H2KXG1RW",
         cost_jpy=2780, stock=1, jan="0000000000000",
         weight_kg=0.03, length=9, width=0.3, height=13,
         images=["https://m.media-amazon.com/images/I/41F0C4gMxaL.jpg"]),
    dict(category="101399", title="P7", desc="P7", parent_sku="B0GQQXJQJC", sku="B0GQQXJQJC",
         cost_jpy=2980, stock=1, jan="0000000000000",
         weight_kg=0.03, length=9, width=0.3, height=13,
         images=["https://m.media-amazon.com/images/I/41n7IpXrcRL.jpg"]),
    dict(category="101399", title="P8", desc="P8", parent_sku="B08LKRYFN6", sku="B08LKRYFN6",
         cost_jpy=1480, stock=1, jan="0000000000000",
         weight_kg=0.05, length=25.7, width=18.2, height=0.1,
         images=["https://m.media-amazon.com/images/I/512+6D1mpAL.jpg"]),
    dict(category="101399", title="P9", desc="P9", parent_sku="B08SVZSCYY", sku="B08SVZSCYY",
         cost_jpy=594, stock=8, jan="0000000000000",
         weight_kg=0.05, length=25, width=18, height=0.1,
         images=["https://m.media-amazon.com/images/I/51u2JM3KRAS.jpg"]),
    dict(category=None, title="P10", desc="P10", parent_sku="B09FQ5RH27", sku="B09FQ5RH27",
         cost_jpy=649, stock=15, jan="0000000000000",
         weight_kg=0.09, length=7, width=23, height=2,
         images=["https://m.media-amazon.com/images/I/71X9AVU-laL._AC_SL1500_.jpg",
                 "https://m.media-amazon.com/images/I/71uxnoSm1rL._AC_SL1500_.jpg",
                 "https://m.media-amazon.com/images/I/41K-Ku18NUL._AC_.jpg",
                 "https://m.media-amazon.com/images/I/31ZAKbFOvIL._AC_.jpg",
                 "https://m.media-amazon.com/images/I/21INRJwwGgL._AC_.jpg"]),
    dict(category="101396", title="P11", desc="P11", parent_sku="B07RWGJHT9", sku="B07RWGJHT9",
         cost_jpy=1780, stock=1, jan="0000000000000",
         weight_kg=0.01, length=12.6, width=8.5, height=0.8,
         images=["https://m.media-amazon.com/images/I/31L2gJmYyiL.jpg"]),
    dict(category="101544", title="P12", desc="P12", parent_sku="B096X73TX8", sku="B096X73TX8",
         cost_jpy=14800, stock=1, jan="0000000000000",
         weight_kg=0.8, length=30, width=25, height=8,
         images=["https://m.media-amazon.com/images/I/61tgdRSWtNS.jpg",
                 "https://m.media-amazon.com/images/I/51qDMmBKyWS.jpg",
                 "https://m.media-amazon.com/images/I/41MqaS9qaDS.jpg",
                 "https://m.media-amazon.com/images/I/41SE7fqGV1S.jpg",
                 "https://m.media-amazon.com/images/I/31guFM9Mt6S.jpg",
                 "https://m.media-amazon.com/images/I/41n-dukOAjS.jpg",
                 "https://m.media-amazon.com/images/I/51enVUqAWQS.jpg"]),
    dict(category="101396", title="P13", desc="P13", parent_sku="B07W4VSXPT", sku="B07W4VSXPT",
         cost_jpy=1475, stock=1, jan="0000000000000",
         weight_kg=0.01, length=12.5, width=8.5, height=0.9,
         images=["https://m.media-amazon.com/images/I/41LD93n2M+L.jpg"]),
    dict(category="101722", title="P14", desc="P14", parent_sku="B0GS8LRZ4Z", sku="B0GS8LRZ4Z",
         cost_jpy=3980, stock=10, jan="0000000000000",
         weight_kg=0.4, length=15, width=10, height=25,
         images=["https://m.media-amazon.com/images/I/41s6hb9zSxL.jpg",
                 "https://m.media-amazon.com/images/I/41VZeIDSqWL.jpg",
                 "https://m.media-amazon.com/images/I/41s25xuFK8L.jpg",
                 "https://m.media-amazon.com/images/I/71OAGyNGBZL._AC_SL1500_.jpg"]),
    dict(category="101396", title="P15", desc="P15", parent_sku="B08592B2Q4", sku="B08592B2Q4",
         cost_jpy=1380, stock=1, jan="0000000000000",
         weight_kg=0.03, length=13.2, width=8.6, height=1.3,
         images=["https://m.media-amazon.com/images/I/7145J2X4WwL.jpg",
                 "https://m.media-amazon.com/images/I/410mdR2YuvL.jpg"]),
    dict(category="101396", title="P16", desc="P16", parent_sku="B08HK56M27", sku="B08HK56M27",
         cost_jpy=3680, stock=1, jan="0000000000000",
         weight_kg=0.05, length=14.8, width=10.9, height=10.8,
         images=["https://m.media-amazon.com/images/I/41KlhVrBZdL._AC_.jpg",
                 "https://m.media-amazon.com/images/I/51A2ZMPbeiL._AC_.jpg",
                 "https://m.media-amazon.com/images/I/41ABLBvm12L._AC_.jpg"]),
]

TITLES = {
    "P1": "Demon Slayer Tanjiro Kamado Infinity Castle Anime Hero PVC Figure JAPAN",
    "P2": "Demon Slayer Season 1 Limited Edition Anime Blu-ray Collector JAPAN",
    "P3": "Demon Slayer Shinobu Kocho Noodle Cup Topper Anime Figure JAPAN",
    "P4": "Demon Slayer Sanemi Shinazugawa Poseable Anime Figure Collector JAPAN",
    "P5": "Demon Slayer Tanjiro Kamado Cross Link Anime Statue Figure JAPAN",
    "P6": "Demon Slayer Shinobu Kocho Acrylic Standee Anime Display JAPAN",
    "P7": "Demon Slayer Giyu Tomioka Acrylic Standee Anime Display JAPAN",
    "P8": "Demon Slayer Movie Key Visual Anime Pencil Board Stationery JAPAN",
    "P9": "Demon Slayer Hashira Group Anime Pencil Board Stationery Gift JAPAN",
    "P10": "Demon Slayer Kids Toothbrush 3 Pack Anime Character Gift JAPAN",
    "P11": "Demon Slayer Tanjiro Kamado Wooden Charm Anime Keychain JAPAN",
    "P12": "Demon Slayer Movie Ufotable Deluxe Art Book Collector Set JAPAN",
    "P13": "Demon Slayer Inosuke Hashibira Wooden Charm Anime Keychain JAPAN",
    "P14": "Demon Slayer Giyu and Tanjiro Big Sitting Plush Set Anime JAPAN",
    "P15": "Demon Slayer Mitsuri Kanroji Acrylic Keychain Anime Charm JAPAN",
    "P16": "Demon Slayer Shinobu Kocho Roly Poly Plush Keychain Anime JAPAN",
}


def build_country_file(country: str):
    src = os.path.join(TEMPLATE_DIR, f"Shopee_mass_upload_template_{country}.xlsx")
    out_path = os.path.join(OUT_DIR, f"Shopee_upload_{country}_{DATE_STR}_asin_batch.xlsx")
    load_fixed_copy(src, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Template"]
    assert ws.max_row == 6, f"{country}: max_row is not 6 before writing ({ws.max_row})"

    row_idx = 7
    for r in ROWS:
        price = calc_local_price(r["cost_jpy"], country)["price_local"]
        parent_asin = r["parent_sku"]
        cover_url = f"{FRAME_BASE_URL}/{parent_asin}_cover.jpg"
        item_images = r["images"][1:]

        if r["category"]:
            ws.cell(row=row_idx, column=1, value=r["category"])
        if r["title"]:
            ws.cell(row=row_idx, column=2, value=TITLES[r["title"]])
        if r["desc"]:
            ws.cell(row=row_idx, column=3, value=DESCRIPTIONS[r["desc"]])
        ws.cell(row=row_idx, column=9, value=parent_asin)  # Parent SKU
        if "vint" in r:
            ws.cell(row=row_idx, column=10, value=r["vint"])
            ws.cell(row=row_idx, column=11, value=r["vname1"])
            ws.cell(row=row_idx, column=12, value=r["opt1"])
        ws.cell(row=row_idx, column=16, value=price)
        ws.cell(row=row_idx, column=17, value=r["stock"])
        ws.cell(row=row_idx, column=18, value=r["sku"])
        ws.cell(row=row_idx, column=21, value=cover_url)
        for i, img_url in enumerate(item_images):
            col = 22 + i
            if col > LAST_ITEM_IMAGE_COL:
                break
            ws.cell(row=row_idx, column=col, value=img_url)
        ws.cell(row=row_idx, column=30, value=r["weight_kg"])
        ws.cell(row=row_idx, column=31, value=r["length"])
        ws.cell(row=row_idx, column=32, value=r["width"])
        ws.cell(row=row_idx, column=33, value=r["height"])
        for col in DELIVERY_COLS[country]:
            ws.cell(row=row_idx, column=col, value="On")
        row_idx += 1

    wb.save(out_path)
    convert_to_shared_strings(out_path)
    print(f"Saved {out_path} ({row_idx - 7} rows)")


if __name__ == "__main__":
    for cc in ["SG", "MY", "TH", "PH"]:
        build_country_file(cc)
