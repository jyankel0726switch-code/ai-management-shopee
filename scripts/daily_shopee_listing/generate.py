# -*- coding: utf-8 -*-
"""Shopee 出品ファイル(SG/PH/MY/TH)を国別テンプレートから生成する。"""
import json, os, sys, zipfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
SKILL = "/home/user/ai-management-shopee/.claude/skills/shopee-listing-upload"
sys.path.insert(0, SKILL)

from pricing_calc import calc_local_price
from listing_copy import COPY, SHIP_BLOCK
from products import PRODUCTS, EST_WEIGHT, EST_DIMS

TEMPLATE = SKILL + "/template/Shopee_mass_upload_template_{cc}.xlsx"
CHANNELS = {
    "SG": ["Doorstep Delivery (Overseas)", "Collection Points (Overseas)", "SPX Express Lockers (Overseas)"],
    "MY": ["Doorstep Delivery - Japan", "SPX Express Lockers (Overseas)"],
    "TH": ["International Express - ส่งจากต่างประเทศ (Japan)"],
    "PH": ["Standard International"],
}
PRICE_BOUNDS = {"SG": (0.10, 999999.00), "MY": (0.10, 1000000000.00),
                "TH": (1, 500000), "PH": (5, 100000)}
NAME_LIMITS = {"SG": (10, 255), "MY": (10, 255), "TH": (20, 255), "PH": (20, 255)}
DESC_LIMITS = {"SG": (20, 3000), "MY": (20, 3000), "TH": (60, 5000), "PH": (20, 3000)}
JAN_PLACEHOLDER = "0000000000000"


def normalize(src, dst):
    """activePane="bottom_left" は OOXML 仕様外で openpyxl が読めないため正規化する。"""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename.startswith("xl/worksheets/"):
                data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
            zout.writestr(it, data)


def build_description(p, cc):
    body = COPY[p["asin"]]["body"]
    jan = p.get("jan") or JAN_PLACEHOLDER
    return f"{body}\n\n{SHIP_BLOCK}\n\nJAN: {jan}"


def main(outdir, date_str):
    images = json.load(open(os.path.join(HERE, "images.json")))
    os.makedirs(outdir, exist_ok=True)
    flags, summary = [], {}

    for cc in ["SG", "PH", "MY", "TH"]:
        rows = [p for p in PRODUCTS if p["cc"] == cc]
        tmp = os.path.join(outdir, f"_norm_{cc}.xlsx")
        normalize(TEMPLATE.format(cc=cc), tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb["Template"]
        assert ws.max_row == 6, f"{cc}: テンプレの max_row={ws.max_row} (6 でなければガイド行が壊れている)"
        H = {(ws.cell(row=3, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}

        lo, hi = PRICE_BOUNDS[cc]
        nlo, nhi = NAME_LIMITS[cc]
        dlo, dhi = DESC_LIMITS[cc]

        for i, p in enumerate(rows):
            r = 7 + i
            put = lambda name, v: ws.cell(row=r, column=H[name], value=v)
            asin = p["asin"]

            title = COPY[asin]["title"]
            assert nlo <= len(title) <= nhi, f"{cc}/{asin}: タイトル長 {len(title)} が {nlo}-{nhi} 外"
            desc = build_description(p, cc)
            assert dlo <= len(desc) <= dhi, f"{cc}/{asin}: 説明文長 {len(desc)} が {dlo}-{dhi} 外"

            price = calc_local_price(p["cost_jpy"], cc)["price_local"]
            assert lo <= price <= hi, f"{cc}/{asin}: 価格 {price} が {lo}-{hi} 外"

            imgs = images.get(asin, [])
            weight = p["weight"] if p["weight"] is not None else EST_WEIGHT
            dims = p["dims"] if p["dims"] is not None else EST_DIMS

            put("Category", p["category"])
            put("Product Name", title)
            put("Product Description", desc)
            put("Parent SKU", asin)
            put("SKU", asin)
            put("Price", round(price, 2))
            put("Stock", p["stock"])
            put("Cover image", imgs[0] if imgs else None)
            for n, url in enumerate(imgs[1:8], start=1):
                put(f"Item Image {n}", url)
            put("Weight", round(float(weight), 2))
            put("Length", dims[0]); put("Width", dims[1]); put("Height", dims[2])
            for ch in CHANNELS[cc]:
                put(ch, "On")

            if not p.get("jan"):
                flags.append((cc, asin, "JANコード未取得 → プレースホルダー 0000000000000"))
            if p["weight"] is None:
                flags.append((cc, asin, f"重量が商品ページ未記載 → 仮値 {EST_WEIGHT}kg"))
            if p["dims"] is None:
                flags.append((cc, asin, f"寸法が商品ページ未記載 → 仮値 {EST_DIMS} cm"))
            if len(imgs) < 7:
                flags.append((cc, asin, f"商品画像 {len(imgs)}枚のみ取得(7枚未満)"))
            flags.append((cc, asin, "Cover image フレーム加工未実施(環境制約) → Amazon元画像URLを使用"))
            if p.get("note"):
                flags.append((cc, asin, p["note"]))

        out = os.path.join(outdir, f"Shopee_upload_{cc}_{date_str}.xlsx")
        wb.save(out)
        os.remove(tmp)
        summary[cc] = (out, len(rows))
        print(f"{cc}: {len(rows)} 行 -> {out}")

    print("\n--- 要確認フラグ ---")
    for cc, asin, msg in flags:
        print(f"{cc} {asin}: {msg}")
    return summary


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
