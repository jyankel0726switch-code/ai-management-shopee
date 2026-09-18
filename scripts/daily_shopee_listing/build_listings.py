"""本日分(2026-09-18 JST)のShopee出品ファイル生成スクリプト"""
import sys
import os

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '.claude', 'skills', 'shopee-listing-upload'))

from xlsx_fix import load_fixed_copy, convert_to_shared_strings
import openpyxl
from pricing_calc import calc_local_price

DATE_STR = "2026-09-18"
# NOTE: SKILL.mdのURL規定は "main" ブランチ固定だが、本セッションは開発ブランチ
# claude/admiring-bardeen-itpmle に画像をpushするため、mainマージ前でも実際に参照できる
# よう、そのブランチを指すURLに変更している(mainへマージ後はmain版に差し替え可能)。
FRAME_BASE_URL = "https://raw.githubusercontent.com/jyankel0726switch-code/ai-management-shopee/claude/admiring-bardeen-itpmle/framed-images"

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), '..', '..', '.claude', 'skills', 'shopee-listing-upload', 'template')
OUT_DIR = os.path.dirname(__file__)

# 配送チャネル列(国ごとに列数・名称が異なる)。すべて On にする。
DELIVERY_COLS = {
    "SG": [34, 35, 36],
    "MY": [34, 35],
    "TH": [34],
    "PH": [34],
}
PREORDER_COL = {"SG": 37, "MY": 36, "TH": 35, "PH": 35}
LAST_ITEM_IMAGE_COL = 29  # Item Image 8 (col 29); Item Image 1 = col22

PRODUCTS = [
    {
        "asin": "B0BGR16QJV",
        "category": "101392",
        "title": "One Piece Shanks Uta Anime Figure Premium Collectible Statue Gift JAPAN",
        "description": """Ever dreamed of holding a piece of the Grand Line right in your own hands?

This premium One Piece collectible figure captures Shanks and Uta in a beautifully sculpted PVC and ABS display piece, painted with rich detail from the iconic Film Red saga. Every line, every color, every expression is built to bring your favorite pirate legend to life on your shelf.

Crafted for real fans who want more than a toy - this is a statement piece. Display it on your desk, your gaming setup, or your collection shelf, and let it spark a conversation every single time a friend walks in.

This is a completed, ready-to-display figure - no assembly needed. Simply unbox it, pose it, and enjoy.

Shipping details: your order is carefully packed and shipped from Japan. Please allow 3-5 business days for us to prepare and dispatch your order, with total delivery time of approximately 1-2 weeks depending on your location.

Do not miss the chance to own this one-of-a-kind piece from one of anime's most legendary sagas.

JAN: 0000000000000""",
        "cost_jpy": 8480,
        "stock": 1,
        "jan": "0000000000000",
        "jan_flag": True,
        "weight_kg": 0.835,
        "weight_flag": False,
        "length": 11.4, "width": 14, "height": 23.9,
        "dim_flag": False,
        "images": [
            "https://m.media-amazon.com/images/I/41tY3stGEbL.jpg",
            "https://m.media-amazon.com/images/I/41TWhadzpxL.jpg",
            "https://m.media-amazon.com/images/I/41SGIZFIyxL.jpg",
            "https://m.media-amazon.com/images/I/51jZ5UR--hL.jpg",
            "https://m.media-amazon.com/images/I/61-Qu8DSe+L.jpg",
            "https://m.media-amazon.com/images/I/61pvGVAjUzL.jpg",
        ],
        "preorder": False,
        "frame_ok": True,
    },
    {
        "asin": "B0FNMWZXGC",
        "category": "101385",
        "title": "Chiikawa Nendoroid Cute Anime Character Poseable Figure Home Decor JAPAN",
        "description": """Isn't it time your desk had something this adorable staring back at you?

This Chiikawa Nendoroid brings the tiny, big-hearted character everyone is obsessed with straight to your desk or shelf. Made from carefully painted plastic, this poseable figure lets you swap expressions and positions, so you can recreate your favorite cute moments any way you like.

Standing at a compact, non-scale size perfect for small spaces, this figure is designed for real fans who want daily doses of comfort and joy. Whether it is sitting on your desk at work, your gaming station, or your bookshelf at home, this little figure has a way of making everything feel a bit softer.

This is a completed, ready-to-display figure - simply unbox, pose, and enjoy immediately.

Shipping details: your order ships from Japan. Please allow 3-5 business days for us to prepare and dispatch your order, with total delivery time of approximately 1-2 weeks depending on your location.

Bring home a little bundle of comfort today.

JAN: 4545784014974""",
        "cost_jpy": 5162,
        "stock": 3,
        "jan": "4545784014974",
        "jan_flag": False,
        "weight_kg": 0.15,
        "weight_flag": True,
        "length": 5, "width": 6, "height": 7,
        "dim_flag": False,
        "images": [
            "https://m.media-amazon.com/images/I/51-xkAPl4XL.jpg",
            "https://m.media-amazon.com/images/I/415rWSvaA2L.jpg",
            "https://m.media-amazon.com/images/I/41OE6IatR6L.jpg",
            "https://m.media-amazon.com/images/I/416qdebyl5L.jpg",
            "https://m.media-amazon.com/images/I/51Sz1+GCyvL.jpg",
            "https://m.media-amazon.com/images/I/51JUaEWsAmL.jpg",
        ],
        "preorder": False,
        "frame_ok": True,
    },
    {
        "asin": "B0G2M5BY1G",
        "category": "101392",
        "title": "Attack on Titan Eren Grandista Anime Hero Statue Collectible JAPAN",
        "description": """What does it feel like to stand face to face with humanity's last hope?

This Attack on Titan Grandista figure brings Eren Yeager to life in dramatic, large-scale detail. Every muscle, every determined expression is sculpted to capture the raw intensity of the story that gripped millions of fans around the world.

Standing at approximately 28cm tall, this figure commands attention on any shelf, desk, or display case. It is built for collectors who want their space to reflect the passion they feel for the series - bold, powerful, and unforgettable.

This is a completed, ready-to-display figure - no building or painting required. Just unbox it and give it the spotlight it deserves.

Shipping details: your order ships from Japan. Please allow 3-5 business days for us to prepare and dispatch your order, with total delivery time of approximately 1-2 weeks depending on your location.

Claim your piece of this legendary story before it is gone.

JAN: 0000000000000""",
        "cost_jpy": 3780,
        "stock": 10,
        "jan": "0000000000000",
        "jan_flag": True,
        "weight_kg": 0.5,
        "weight_flag": True,
        "length": 12, "width": 18, "height": 28,
        "dim_flag": False,
        "images": [
            "https://m.media-amazon.com/images/I/31NfjWM3jLL.jpg",
            "https://m.media-amazon.com/images/I/51XEPtJC-VL.jpg",
        ],
        "preorder": False,
        "frame_ok": True,
    },
    {
        "asin": "B0CPBBGZSB",
        "category": "101385",
        "title": "Haikyu Bokuto Volleyball Anime Hero Posing Figure Sports Fan Gift JAPAN",
        "description": """Remember the thrill of watching your favorite player spike the winning point?

This Haikyu posing figure brings Bokuto to life in a dynamic action pose, capturing the energy and passion that made this volleyball story a fan favorite everywhere. Every detail, from his expression to his stance, is designed to feel like a frozen moment from the court.

Compact and lightweight, this figure is perfect for a desk, a shelf, or a display case, and makes a heartfelt gift for anyone who loves volleyball or great sports stories. It is a small piece that carries a big feeling.

This is a completed, ready-to-display figure - no assembly needed, simply unbox and display.

Shipping details: your order ships from Japan. Please allow 3-5 business days for us to prepare and dispatch your order, with total delivery time of approximately 1-2 weeks depending on your location.

Bring the excitement of the court home today.

JAN: 0000000000000""",
        "cost_jpy": 4300,
        "stock": 1,
        "jan": "0000000000000",
        "jan_flag": True,
        "weight_kg": 0.35,
        "weight_flag": False,
        "length": 5.1, "width": 8.1, "height": 16,
        "dim_flag": True,
        "images": [
            "https://m.media-amazon.com/images/I/61HQ3MvYV9L.jpg",
            "https://m.media-amazon.com/images/I/31n16yjk-7L.jpg",
            "https://m.media-amazon.com/images/I/51W+e86SG1L.jpg",
            "https://m.media-amazon.com/images/I/51qf6mg6zhL.jpg",
            "https://m.media-amazon.com/images/I/41h7Xh822KL.jpg",
            "https://m.media-amazon.com/images/I/41M1Og0lTOL.jpg",
        ],
        "preorder": False,
        "frame_ok": True,
    },
]


def build_country_file(country: str):
    src = os.path.join(TEMPLATE_DIR, f"Shopee_mass_upload_template_{country}.xlsx")
    out_path = os.path.join(OUT_DIR, f"Shopee_upload_{country}_{DATE_STR}.xlsx")
    load_fixed_copy(src, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Template"]
    assert ws.max_row == 6, f"{country}: max_row is not 6 before writing ({ws.max_row})"

    row_idx = 7
    for p in PRODUCTS:
        price = calc_local_price(p["cost_jpy"], country)["price_local"]
        cover_url = f"{FRAME_BASE_URL}/{p['asin']}_cover.jpg" if p["frame_ok"] else p["images"][0]
        item_images = p["images"][1:]  # 2枚目以降をItem Image1〜

        ws.cell(row=row_idx, column=1, value=p["category"])          # Category
        ws.cell(row=row_idx, column=2, value=p["title"])              # Product Name
        ws.cell(row=row_idx, column=3, value=p["description"])        # Product Description
        ws.cell(row=row_idx, column=9, value=p["asin"])                # Parent SKU
        ws.cell(row=row_idx, column=16, value=price)                   # Price
        ws.cell(row=row_idx, column=17, value=p["stock"])               # Stock
        ws.cell(row=row_idx, column=18, value=p["asin"])                # SKU
        ws.cell(row=row_idx, column=21, value=cover_url)                # Cover image

        for i, img_url in enumerate(item_images):
            col = 22 + i  # Item Image 1 = col22
            if col > LAST_ITEM_IMAGE_COL:
                break
            ws.cell(row=row_idx, column=col, value=img_url)

        ws.cell(row=row_idx, column=30, value=p["weight_kg"])           # Weight (kg)
        ws.cell(row=row_idx, column=31, value=p["length"])              # Length
        ws.cell(row=row_idx, column=32, value=p["width"])               # Width
        ws.cell(row=row_idx, column=33, value=p["height"])              # Height

        for col in DELIVERY_COLS[country]:
            ws.cell(row=row_idx, column=col, value="On")

        if p["preorder"]:
            ws.cell(row=row_idx, column=PREORDER_COL[country], value=10)

        row_idx += 1

    wb.save(out_path)
    convert_to_shared_strings(out_path)
    print(f"Saved {out_path} ({row_idx - 7} rows)")


if __name__ == "__main__":
    for cc in ["SG", "MY", "TH", "PH"]:
        build_country_file(cc)
